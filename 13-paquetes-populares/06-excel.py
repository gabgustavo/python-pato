import openpyxl

sheets = openpyxl.load_workbook('13-paquetes-populares/planilla.xlsx')

# print(sheets.sheetnames)
# print(sheets['Hoja 1'])
hoja = sheets.active

# print(hoja)

sheets.create_sheet('Mi Nueva data')
# print(sheets.sheetnames)

miHoja = sheets['Mi Nueva data']
miHoja.title = 'Mi Nueva Data'


print(f"""
      Filas: {hoja.max_row}
      Columnas: {hoja.max_column}
""")

celda = hoja["A1"]
print(celda)
print(celda.value)
celda.value = "Nombres"

celda2 = hoja.cell(row=2, column=1)
print(celda2.value)

print(f"""
    en celda 2 hay en valor:      {celda2.value}
    en celda 2 hay en row:        {celda2.row}
    en celda 2 hay en column:     {celda2.column}
    en celda 2 hay en coordinate: {celda2.coordinate}
""")

print("="*30)
for fila in range(1, hoja.max_row + 1):
    for columna in range(1, hoja.max_column + 1):
        celda = hoja.cell(row=fila, column=columna)
        print(f"{celda.coordinate}: {celda.value}")

columna = hoja["A"]
fila = hoja["1"]
print(columna)
print(fila)

hoja.append([1, 2, 3])
print(hoja.rows)

# hoja.delete_cols()
hoja.delete_rows(1, 1)

sheets.save('13-paquetes-populares/nuevo_documento_excel.xlsx')
